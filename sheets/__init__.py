"""
sheets/__init__.py
==================
Shared helpers for all jadual mappers and the engine.

The plugin registry in `src/engine.py` scans this package for modules
named `jadual_*.py`, so this file itself is never loaded as a mapper.
"""
import math
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ==========================================
# RICH TEXT  (optional — requires openpyxl >= 3.1)
# ==========================================
try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    RICH_TEXT_AVAILABLE = True
except ImportError:
    RICH_TEXT_AVAILABLE = False
    print("[notice] Rich text not available. Install openpyxl >= 3.1")


# ==========================================
# STYLE CONSTANTS
# ==========================================
FONT_NAME = "Century Gothic"
FONT_SIZE = 11

CELL_FONT       = Font(name=FONT_NAME, size=FONT_SIZE)
CELL_ALIGN      = Alignment(horizontal="right", vertical="top",   wrap_text=False)
BILINGUAL_ALIGN = Alignment(horizontal="left",  vertical="center", wrap_text=True)


# ==========================================
# VALUE FORMATTING
# ==========================================
def _format_value(val, decimals: int | None) -> int | float | str:
    """
    Universal rules:
      • None / "" / NaN / "n.a"  →  "n.a"
      • non-numeric text         →  unchanged
      • whole number             →  int
      • decimal number           →  rounded to `decimals` places (half-up)
    `decimals = None` means: don't round, but still integerize whole numbers.
    """
    if val is None or val == "" or val in ("n.a", "N.A", "na"):
        return "n.a"

    if isinstance(val, float) and math.isnan(val):
        return "n.a"

    try:
        num = Decimal(str(val))
    except (InvalidOperation, ValueError, TypeError):
        return val

    if decimals is None:
        return int(num) if num == num.to_integral_value() else float(num)

    quant   = Decimal(1).scaleb(-decimals)
    rounded = num.quantize(quant, rounding=ROUND_HALF_UP)
    return int(rounded) if rounded == rounded.to_integral_value() else float(rounded)


# ==========================================
# CELL WRITER
# ==========================================
def write_cell(ws, row: int, col: int, value, decimals: int | None = 1) -> None:
    """Write value + standard font/alignment + universal rounding."""
    cell = ws.cell(row=row, column=col)
    cell.value     = _format_value(value, decimals)
    cell.font      = CELL_FONT
    cell.alignment = CELL_ALIGN


# ==========================================
# BILINGUAL RICH-TEXT FORMATTING
# ==========================================
def extract_plain_text(value) -> str:
    """Flatten a CellRichText (or any value) into a plain string."""
    if value is None:
        return ""
    if RICH_TEXT_AVAILABLE and isinstance(value, CellRichText):
        parts = []
        for item in value:
            if isinstance(item, TextBlock):
                parts.append(item.text)
            else:
                parts.append(str(item))
        return "".join(parts)
    return str(value)


def apply_bilingual_rich_formatting(ws, row: int, col: int,
                                    font_size: int | None = None,
                                    debug: bool = False) -> None:
    """
    Line 1 of a cell → bold;  lines 2+ → italic.
    Used on the header labels in column B ("MAKLUMAT ASAS / BASIC INFORMATION")
    and column C ("Tahun / Year") of every page.
    """
    cell = ws.cell(row=row, column=col)
    cell_ref = f"{get_column_letter(col)}{row}"

    if not RICH_TEXT_AVAILABLE:
        return

    original = cell.value
    if original is None:
        if debug:
            print(f"    [{cell_ref}] empty; skipped")
        return

    text = extract_plain_text(original)
    if not text.strip():
        return

    lines = text.split("\n")
    if debug:
        preview = " | ".join(repr(l) for l in lines)
        print(f"    [{cell_ref}] lines={len(lines)}: {preview}")

    # Single line → leave as plain text, only ensure wrap
    if len(lines) < 2:
        cell.value = text
        if not cell.alignment or not cell.alignment.wrap_text:
            cell.alignment = BILINGUAL_ALIGN
        return

    size_str    = str(font_size if font_size is not None else FONT_SIZE)
    bold_font   = InlineFont(rFont=FONT_NAME, sz=size_str, b=True)
    italic_font = InlineFont(rFont=FONT_NAME, sz=size_str, i=True)

    blocks = []
    first_text = lines[0] + "\n"
    blocks.append(TextBlock(bold_font, first_text))

    for i in range(1, len(lines)):
        is_last   = (i == len(lines) - 1)
        line_text = lines[i] if is_last else lines[i] + "\n"
        blocks.append(TextBlock(italic_font, line_text))

    cell.value = CellRichText(*blocks)

    existing = cell.alignment
    if not (existing and existing.wrap_text):
        cell.alignment = BILINGUAL_ALIGN